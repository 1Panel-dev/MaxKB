"""Excel sheet boundaries survive extraction and workflow document splitting."""

import io
from unittest.mock import MagicMock, patch

import openpyxl
import xlwt
from django.test import SimpleTestCase

from application.workflow.common import WorkflowType
from application.workflow.nodes.document_extract_node.document_extract_node import DocumentExtractNode
from application.workflow.nodes.document_split_node.document_split_node import DocumentSplitNode


class WorkflowExcelSplitTests(SimpleTestCase):
    file_id = "00000000-0000-0000-0000-000000000001"
    source_meta = {"source_type": "local", "source_scope": "local-source"}

    def workbook_bytes(self, suffix, sheet_names):
        workbook = xlwt.Workbook() if suffix.lower() == ".xls" else openpyxl.Workbook()
        if suffix.lower() != ".xls":
            workbook.remove(workbook.active)
        for name in sheet_names:
            rows = [
                ["分段标题", "分段内容", "问题"],
                [f"{name}标题", f"{name}回答", f"{name}问题"],
                [f"{name}标题二", f"{name}回答二", f"{name}问题二"],
            ]
            if suffix.lower() == ".xls":
                sheet = workbook.add_sheet(name)
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        sheet.write(row_index, column_index, value)
            else:
                sheet = workbook.create_sheet(name)
                for row in rows:
                    sheet.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def extract(self, suffix, sheet_names):
        file = MagicMock(id=self.file_id)
        file.get_bytes.return_value = self.workbook_bytes(suffix, sheet_names)
        node = DocumentExtractNode.__new__(DocumentExtractNode)
        node.get_parameters = MagicMock(return_value={"document_list": ["source", "file_list"]})
        node.get_workflow_parameters = MagicMock(return_value={"knowledge_id": "knowledge-id"})
        node.get_workflow_type = MagicMock(return_value=WorkflowType.KNOWLEDGE)
        node.workflow_manage = MagicMock()
        node.workflow_manage.get_reference_field.return_value = [
            {"file_id": self.file_id, "name": f"knowledge{suffix}", "meta": self.source_meta}
        ]
        context = {}
        node.write_context = context.__setitem__
        with (
            patch("application.workflow.nodes.document_extract_node.document_extract_node.QuerySet") as query_set,
            patch("common.handle.impl.table.xlsx_parse_table_handle.xlsx_embed_cells_images", return_value={}),
        ):
            query_set.return_value.filter.return_value.first.return_value = file
            node.execute()
        return context["document_list"]

    def split(self, documents, **parameters):
        node = DocumentSplitNode.__new__(DocumentSplitNode)
        node.get_parameters = MagicMock(return_value={"document_list": ["extract", "document_list"], **parameters})
        node.get_workflow_parameters = MagicMock(return_value={"knowledge_id": "knowledge-id"})
        node.get_workflow_type = MagicMock(return_value=WorkflowType.KNOWLEDGE)
        node.get_reference_content = MagicMock(return_value=documents)
        context = {}
        node.write_context = context.__setitem__
        node.execute()
        return context["paragraph_list"]

    def test_smart_split_keeps_each_sheet_and_its_source_file_separate(self):
        for suffix in (".xlsx", ".XLSX", ".xls"):
            with self.subTest(suffix=suffix):
                documents = self.split(self.extract(suffix, ["Sheet1", "产品说明"]))
                self.assertEqual([document["name"] for document in documents], ["Sheet1", "产品说明"])
                for document, other_name in zip(documents, ["产品说明", "Sheet1"]):
                    self.assertEqual(document["source_file_id"], self.file_id)
                    self.assertEqual(
                        document["meta"],
                        {**self.source_meta, "source_file_id": self.file_id, "knowledge_id": "knowledge-id"},
                    )
                    self.assertEqual(len(document["paragraphs"]), 1)
                    paragraph = document["paragraphs"][0]
                    self.assertEqual(paragraph["title"].strip(), document["name"])
                    self.assertIn(f"{document['name']}回答", paragraph["content"])
                    self.assertNotIn(other_name, paragraph["content"])
                    self.assertTrue(paragraph["is_active"])
                    self.assertTrue(paragraph["chunks"])

    def test_custom_patterns_cannot_combine_different_sheets(self):
        documents = self.split(
            self.extract(".xlsx", ["Sheet1", "产品说明"]),
            split_strategy="custom",
            patterns=["NEVER_MATCH_THIS_SEPARATOR"],
        )
        self.assertEqual([document["name"] for document in documents], ["Sheet1", "产品说明"])
        self.assertNotIn("产品说明", documents[0]["paragraphs"][0]["content"])
        self.assertNotIn("Sheet1", documents[1]["paragraphs"][0]["content"])

    def test_length_limit_can_create_multiple_paragraphs_within_each_sheet(self):
        documents = self.split(self.extract(".xlsx", ["Sheet1", "产品说明"]), limit=60, chunk_size=20)
        self.assertEqual([document["name"] for document in documents], ["Sheet1", "产品说明"])
        self.assertTrue(all(len(document["paragraphs"]) > 1 for document in documents))

    def test_qa_keeps_sheet_documents_and_row_questions(self):
        for suffix in (".xlsx", ".xls"):
            with self.subTest(suffix=suffix):
                documents = self.split(self.extract(suffix, ["Sheet1", "产品说明"]), split_strategy="qa")
                self.assertEqual([document["name"] for document in documents], ["Sheet1", "产品说明"])
                for document in documents:
                    self.assertEqual(len(document["paragraphs"]), 2)
                    paragraph = document["paragraphs"][0]
                    self.assertEqual(paragraph["content"], f"{document['name']}回答")
                    self.assertEqual(paragraph["problem_list"], [f"{document['name']}问题"])

    def test_single_sheet_name_matches_qa_including_filename_like_titles(self):
        for name in ("Sheet1", "http说明"):
            with self.subTest(sheet=name):
                extracted = self.extract(".xlsx", [name])
                smart = self.split(extracted)
                qa = self.split(extracted, split_strategy="qa")
                self.assertEqual([document["name"] for document in smart], [name])
                self.assertEqual([document["name"] for document in smart], [document["name"] for document in qa])

    def test_non_excel_headings_stay_in_one_document(self):
        for name in ("notes.md", "notes.txt", "notes.docx"):
            with self.subTest(name=name):
                documents = self.split([{"name": name, "content": "## 第一章\n正文一\n\n## 第二章\n正文二"}])
                self.assertEqual([document["name"] for document in documents], [name])
                self.assertEqual(len(documents[0]["paragraphs"]), 2)

    def test_excel_without_sheet_headings_preserves_original_name(self):
        documents = self.split([{"name": "empty.xlsx", "content": ""}])
        self.assertEqual([document["name"] for document in documents], ["empty.xlsx"])
        self.assertEqual(documents[0]["paragraphs"], [])
