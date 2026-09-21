"""Archive round trips exercise real JSON, XLSX and ZIP with isolated model/storage writes."""

import copy
import io
import json
from contextlib import ExitStack
from unittest.mock import MagicMock, patch
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from django.test import SimpleTestCase
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from tools.models import ToolWorkflow

from knowledge.models import (
    Document,
    DocumentResourceType,
    DocumentTag,
    File,
    FileSourceType,
    Knowledge,
    KnowledgeType,
    KnowledgeWorkflow,
    Paragraph,
    ParagraphAsset,
    Problem,
    ProblemParagraphMapping,
    Tag,
    Termbase,
)
from knowledge.serializers import knowledge as serializers
from knowledge.serializers import knowledge_workflow as workflow_serializers
from knowledge.services import knowledge_archive as archive_service
from knowledge.services.document_strategy import normalize_document_strategy


def uid(number):
    return UUID(f"00000000-0000-0000-0000-{number:012d}")


class KnowledgeArchiveTests(SimpleTestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.now = timezone.now()
        self.knowledge = Knowledge(
            id=uid(1),
            name="archive",
            desc="description",
            type=KnowledgeType.WEB,
            user_id=uid(2),
            workspace_id="source",
            folder_id="source",
            external_service={"enabled": True, "authentication": True},
            meta={
                "source_url": "https://example.com",
                "doc_strategy": {"split": {"child_length": 128}},
                "sync_setting": {
                    "enabled": True,
                    "schedule_type": "cron",
                    "cron_expression": "0 2 * * *",
                    "sync_type": "incremental",
                },
            },
        )
        self.file = File(id=uid(3), file_name="风景.png", source_id=str(uid(4)), source_type=FileSourceType.DOCUMENT)
        self.file_bytes = b"image-payload"
        self.stack.enter_context(patch.object(self.file, "get_bytes", return_value=self.file_bytes))
        self.document = Document(
            id=uid(4),
            knowledge_id=self.knowledge.id,
            name="original/name:with special characters",
            resource_type=DocumentResourceType.IMAGE,
            type=KnowledgeType.BASE,
            char_length=6,
            meta={"source_file_id": str(self.file.id), "image_file_id": str(self.file.id), "allow_download": True},
            doc_strategy=normalize_document_strategy(
                {"split": {"child_length": 128}, "index": {"title_as_question": True}}
            ),
            hit_num=9,
            last_hit_time=self.now,
            sync_version=3,
            source_hash="document-hash",
        )
        self.paragraph = Paragraph(
            id=uid(5),
            document_id=self.document.id,
            knowledge_id=self.knowledge.id,
            title="image",
            content=f"![image](./oss/file/{self.file.id})",
            chunks=["child one", "child two"],
            content_schema=[{"type": "image", "file_id": str(self.file.id), "description": "visual description"}],
            source_snapshot={"content": f"![image](./oss/file/{self.file.id})"},
            origin="synced",
            source_key="remote:one",
            source_hash="paragraph-hash",
            hit_num=8,
            last_hit_time=self.now,
        )
        self.asset = ParagraphAsset(
            id=uid(6),
            document_id=self.document.id,
            paragraph_id=self.paragraph.id,
            knowledge_id=self.knowledge.id,
            file_id=self.file.id,
            caption="caption",
            ocr_text="OCR",
            description="visual description",
            hit_num=7,
            last_hit_time=self.now,
            source_asset_key="remote:image:1",
            process_status="success",
        )
        self.problem = Problem(
            id=uid(7), knowledge_id=self.knowledge.id, content="question", hit_num=6, last_hit_time=self.now
        )
        self.mapping = ProblemParagraphMapping(
            id=uid(8),
            knowledge_id=self.knowledge.id,
            document_id=self.document.id,
            paragraph_id=self.paragraph.id,
            problem_id=self.problem.id,
            meta={"kind": "title"},
        )
        self.tag = Tag(id=uid(9), knowledge_id=self.knowledge.id, key="key:one", value="a|b")
        self.rows = {
            Knowledge: [self.knowledge],
            Document: [self.document],
            Paragraph: [self.paragraph],
            ParagraphAsset: [self.asset],
            Problem: [self.problem],
            ProblemParagraphMapping: [self.mapping],
            Tag: [self.tag],
            File: [self.file],
            Termbase: [],
            ToolWorkflow: [],
            DocumentTag: [],
        }
        self.queries = {}
        self.created = {}
        for model, rows in self.rows.items():
            query = MagicMock()
            query.filter.return_value = query
            query.order_by.return_value = query
            query.__iter__.side_effect = lambda rows=rows: iter(rows)
            query.first.side_effect = lambda rows=rows: rows[0] if rows else None
            query.bulk_create.side_effect = lambda values, model=model, **kwargs: self.created.setdefault(
                model, []
            ).extend(values)
            query.values.side_effect = lambda *fields, rows=rows: [
                {field: getattr(row, field) for field in fields} for row in rows
            ]
            query.values_list.return_value = []
            self.queries[model] = query
        self.queries[DocumentTag].values.return_value = [{"document_id": self.document.id, "tag_id": self.tag.id}]
        self.queries[DocumentTag].values.side_effect = None
        for module in (archive_service, serializers):
            self.stack.enter_context(patch.object(module, "QuerySet", side_effect=lambda model: self.queries[model]))
        self.file_writes = []

        def save_file(file, bytea):
            self.file_writes.append((file, bytea))

        self.stack.enter_context(patch.object(File, "save", save_file))
        self.saved_knowledge = []

        def save_knowledge(knowledge, **kwargs):
            self.saved_knowledge.append(knowledge)

        self.stack.enter_context(patch.object(Knowledge, "save", save_knowledge))
        self.stack.enter_context(patch.object(serializers, "UserResourcePermissionSerializer"))
        self.stack.enter_context(patch.object(serializers, "update_resource_mapping_by_knowledge"))
        self.on_commit = self.stack.enter_context(patch.object(serializers.transaction, "on_commit"))

    def export(self, with_source_file=False):
        paragraph_rows = [
            {
                **archive_service._record(self.paragraph, archive_service.PARAGRAPH_FIELDS),
                "document_name": self.document.name,
            }
        ]
        with (
            patch.object(
                serializers,
                "native_search",
                side_effect=[paragraph_rows, [{"content": self.problem.content, "paragraph_id": self.paragraph.id}]],
            ),
            patch.object(serializers, "write_image"),
        ):
            serializer = serializers.KnowledgeSerializer.Operate(
                data={"knowledge_id": str(self.knowledge.id), "user_id": str(uid(2)), "workspace_id": "source"}
            )
            serializers.serializers.Serializer.is_valid(serializer, raise_exception=True)
            response = serializer.export_knowledge(with_source_file=with_source_file, with_valid=False)
        return response.content

    def import_archive(self, payload):
        serializer = serializers.KnowledgeSerializer.ImportKnowledge(
            data={"workspace_id": "target", "folder_id": "target-folder", "user_id": str(uid(10))}
        )
        serializer.is_valid(raise_exception=True)
        return serializers.KnowledgeSerializer.ImportKnowledge.import_knowledge.__wrapped__(
            serializer, io.BytesIO(payload), with_valid=False
        )

    def test_real_zip_export_import_preserves_new_features_and_remaps_all_relations(self):
        payload = self.export()
        with ZipFile(io.BytesIO(payload)) as archive:
            data = json.loads(archive.read("knowledge.json"))
            self.assertIn("knowledge.xlsx", archive.namelist())
            self.assertEqual(data["archive_version"], 2)
            self.assertEqual(data["source_file_list"], [])
            self.assertEqual(archive.read(f"oss/file/{self.file.id}"), self.file_bytes)
        result = self.import_archive(payload)
        imported = self.saved_knowledge[0]
        self.assertEqual(imported.external_service, self.knowledge.external_service)
        self.assertEqual(imported.meta, self.knowledge.meta)
        self.assertEqual(imported.workspace_id, "target")
        self.assertIsNone(imported.embedding_model_id)
        document = self.created[Document][0]
        paragraph = self.created[Paragraph][0]
        asset = self.created[ParagraphAsset][0]
        problem = self.created[Problem][0]
        mapping = self.created[ProblemParagraphMapping][0]
        file, contents = self.file_writes[0]
        self.assertNotEqual(str(document.id), str(self.document.id))
        self.assertEqual(str(document.knowledge_id), result["knowledge_id"])
        self.assertEqual(str(document.user_id), str(uid(10)))
        self.assertEqual(document.resource_type, "image")
        self.assertEqual(document.name, self.document.name)
        self.assertEqual(document.doc_strategy, self.document.doc_strategy)
        self.assertEqual(document.sync_version, 3)
        self.assertEqual(document.hit_num, 9)
        self.assertEqual(document.last_hit_time, self.now)
        self.assertEqual(document.meta["image_file_id"], str(file.id))
        self.assertEqual(document.meta["source_file_id"], str(file.id))
        self.assertEqual(contents, self.file_bytes)
        self.assertEqual(file.file_name, "风景.png")
        self.assertEqual(file.source_id, str(document.id))
        self.assertEqual(paragraph.document_id, str(document.id))
        self.assertEqual(paragraph.chunks, self.paragraph.chunks)
        self.assertEqual(paragraph.content_schema[0]["file_id"], str(file.id))
        self.assertIn(str(file.id), paragraph.content)
        self.assertIn(str(file.id), paragraph.source_snapshot["content"])
        self.assertEqual(paragraph.hit_num, 8)
        self.assertEqual(asset.paragraph_id, str(paragraph.id))
        self.assertEqual(asset.file_id, str(file.id))
        self.assertEqual(asset.description, "visual description")
        self.assertEqual(asset.ocr_text, "OCR")
        self.assertEqual(asset.hit_num, 7)
        self.assertEqual(problem.hit_num, 6)
        self.assertEqual(mapping.problem_id, str(problem.id))
        self.assertEqual(mapping.paragraph_id, str(paragraph.id))
        self.assertEqual(self.created[Tag][0].value, "a|b")
        self.assertEqual(str(self.created[DocumentTag][0].tag_id), str(self.created[Tag][0].id))
        callback = self.on_commit.call_args.args[0]
        self.assertEqual(callback.args, (result["knowledge_id"],))

    def test_original_source_file_is_restored_once_when_also_an_image_asset(self):
        self.import_archive(self.export(with_source_file=True))
        self.assertEqual(len(self.file_writes), 1)
        self.assertEqual(self.file_writes[0][1], self.file_bytes)

    def test_optional_original_file_does_not_replace_image_source_with_another_attachment(self):
        attachment = File(
            id=uid(11), file_name="attachment.pdf", source_id=str(self.document.id), source_type=FileSourceType.DOCUMENT
        )
        self.rows[File].append(attachment)
        with patch.object(attachment, "get_bytes", return_value=b"attachment"):
            self.import_archive(self.export(with_source_file=True))
        source = self.created[Document][0].meta["source_file_id"]
        self.assertEqual(next(content for file, content in self.file_writes if str(file.id) == source), self.file_bytes)

    def test_workflow_bundle_preserves_default_models_and_sync_inputs(self):
        self.knowledge.type = KnowledgeType.WORKFLOW
        self.knowledge.meta["workflow_sync_input"] = {"data_source": {"query": "saved input"}}
        workflow = KnowledgeWorkflow(
            knowledge_id=self.knowledge.id,
            workspace_id="source",
            work_flow={"nodes": [], "edges": []},
            default_model_setting={"llm": {"model_id": str(uid(20))}},
        )
        query = MagicMock()
        query.filter.return_value.first.return_value = workflow
        self.queries[KnowledgeWorkflow] = query
        payload = self.export()
        import_class = workflow_serializers.KnowledgeWorkflowSerializer.Import
        with (
            patch.object(import_class, "import_", import_class.import_.__wrapped__),
            patch.object(KnowledgeWorkflow, "objects") as manager,
            patch.object(workflow_serializers, "update_resource_mapping_by_knowledge"),
        ):
            self.import_archive(payload)
        defaults = manager.filter.return_value.update_or_create.call_args.kwargs["defaults"]
        self.assertEqual(defaults["default_model_setting"], workflow.default_model_setting)
        self.assertEqual(defaults["work_flow"], workflow.work_flow)
        self.assertEqual(
            self.saved_knowledge[0].meta["workflow_sync_input"], self.knowledge.meta["workflow_sync_input"]
        )

    def test_long_content_and_empty_documents_survive_excel_limits(self):
        self.paragraph.content = "长" * 40000
        empty = Document(id=uid(11), knowledge_id=self.knowledge.id, name="empty", char_length=0)
        self.rows[Document].append(empty)
        self.import_archive(self.export())
        self.assertEqual(self.created[Paragraph][0].content, self.paragraph.content)
        self.assertEqual(len(self.created[Document]), 2)
        self.assertEqual(self.created[Document][1].name, "empty")

    def test_lark_export_retains_schedule_but_omits_connection_credentials(self):
        self.knowledge.type = KnowledgeType.LARK
        self.knowledge.meta.update(app_id="connection-id", app_secret="fixture-secret", folder_token="remote")
        with ZipFile(io.BytesIO(self.export())) as archive:
            meta = json.loads(archive.read("knowledge.json"))["meta"]
        self.assertEqual(meta["sync_setting"], self.knowledge.meta["sync_setting"])
        self.assertEqual(set(meta), {"sync_setting", "doc_strategy"})

    def test_unlinked_problems_and_paragraph_anchors_survive(self):
        anchor = Paragraph(id=uid(11), document_id=self.document.id, knowledge_id=self.knowledge.id, content="anchor")
        self.rows[Paragraph].append(anchor)
        self.paragraph.anchor_paragraph_id = anchor.id
        self.rows[Problem].append(Problem(id=uid(12), knowledge_id=self.knowledge.id, content="unlinked", hit_num=4))
        self.import_archive(self.export())
        paragraphs = self.created[Paragraph]
        self.assertEqual(paragraphs[0].anchor_paragraph_id, str(paragraphs[1].id))
        self.assertEqual(self.created[Problem][1].content, "unlinked")
        self.assertEqual(self.created[Problem][1].hit_num, 4)

    def test_ordinary_source_files_remain_optional(self):
        self.document.resource_type = DocumentResourceType.DOCUMENT
        self.document.meta.pop("image_file_id")
        # The original binary is independent from inline image assets.
        self.document.meta["source_file_id"] = str(uid(20))
        with ZipFile(io.BytesIO(self.export())) as archive:
            data = json.loads(archive.read("knowledge.json"))
        self.assertNotIn("source_file_id", data["resources"]["documents"][0]["meta"])

    def test_missing_file_prevents_incomplete_export(self):
        self.rows[File].clear()
        with self.assertRaises(ValidationError):
            self.export()

    def test_foreign_file_reference_cannot_export_another_knowledge_file(self):
        self.file.source_id = str(uid(99))
        with self.assertRaises(ValidationError):
            self.export()

    def test_legacy_auth_marker_and_disabled_schedule_are_preserved(self):
        self.knowledge.external_service = {}
        self.knowledge.meta["sync_setting"]["enabled"] = False
        self.import_archive(self.export())
        self.assertEqual(self.saved_knowledge[0].external_service, {})
        self.assertFalse(self.saved_knowledge[0].meta["sync_setting"]["enabled"])
        self.on_commit.assert_not_called()

    def test_invalid_or_missing_resource_references_are_rejected_before_writes(self):
        payload = self.export()
        with ZipFile(io.BytesIO(payload)) as archive:
            data = json.loads(archive.read("knowledge.json"))
            for change in (
                lambda value: value.update(archive_version=999),
                lambda value: value["resources"]["assets"][0].update(file_id=str(uid(99))),
                lambda value: value["resources"]["paragraphs"][0].update(document_id=str(uid(99))),
                lambda value: value["resources"]["files"][0].update(zip_path="missing.png"),
                lambda value: value.update(external_service={"authentication": "false"}),
                lambda value: value["meta"].update(sync_setting={"schedule_type": "unknown"}),
                lambda value: value["meta"].update(sync_setting={"enabled": "false"}),
            ):
                invalid = copy.deepcopy(data)
                change(invalid)
                with self.assertRaises(ValidationError):
                    archive_service.validate_archive(invalid, archive)
        self.assertEqual(self.saved_knowledge, [])
        self.assertEqual(self.file_writes, [])

    def test_legacy_excel_archive_still_imports_without_new_fields(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "legacy"
        workbook.active.append(["title", "content", "problem"])
        workbook.active.append(["heading", "old content", ""])
        xlsx = io.BytesIO()
        workbook.save(xlsx)
        payload = io.BytesIO()
        with ZipFile(payload, "w", ZIP_DEFLATED) as archive:
            archive.writestr("knowledge.json", json.dumps({"name": "old knowledge"}))
            archive.writestr("knowledge.xlsx", xlsx.getvalue())
        with patch.object(serializers, "ProblemParagraphManage") as problems:
            problems.return_value.to_problem_model_list.return_value = ([], [])
            self.import_archive(payload.getvalue())
        self.assertEqual(self.created[Document][0].resource_type, "document")
        self.assertEqual(self.created[Paragraph][0].content, "old content")
        self.assertEqual(self.saved_knowledge[0].external_service, {"enabled": False, "authentication": False})
        self.on_commit.assert_not_called()

    def test_file_and_workflow_input_references_are_rebound(self):
        self.knowledge.meta["workflow_sync_input"] = {"data_source": {"file_id": str(self.file.id)}}
        self.import_archive(self.export())
        self.assertEqual(
            self.saved_knowledge[0].meta["workflow_sync_input"]["data_source"]["file_id"],
            str(self.file_writes[0][0].id),
        )

    def test_file_links_use_destination_relative_paths(self):
        for prefix in (".", "", "/custom/chat", "https://old.example/custom/chat"):
            with self.subTest(prefix=prefix):
                value = f"![image]({prefix}/oss/file/{uid(3)})"
                self.assertEqual(
                    archive_service.remap_references(value, {str(uid(3)): uid(30)}),
                    f"![image](./oss/file/{uid(30)})",
                )
