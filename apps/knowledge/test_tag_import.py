import io
import json
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import quote
from uuid import UUID

import openpyxl
import xlwt
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.urls import path, resolve
from drf_spectacular.generators import SchemaGenerator
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIRequestFactory, force_authenticate

from common.auth import TokenAuth
from common.auth.constants.permission_constants import PermissionConstants
from common.exception.app_exception import AppApiException, AppUnauthorizedFailed
from knowledge.api.tag_import import TagImportAPI, TagTemplateExportAPI
from knowledge.serializers.tag_import import (
    TAG_TEMPLATE_CONTENT_TYPE,
    TAG_TEMPLATE_PATH,
    TagImportSerializers,
    parse_tag_file,
)
from knowledge.views.tag_import import KnowledgeTagImportView, KnowledgeTagTemplateView


KNOWLEDGE_ID = "00000000-0000-0000-0000-000000000001"
USER_ID = "00000000-0000-0000-0000-000000000002"


def tag_file(rows, extension="xlsx", extra_sheets=None):
    buffer = io.BytesIO()
    sheets = {"Sheet1": rows, **(extra_sheets or {})}
    if extension == "xlsx":
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for name, sheet_rows in sheets.items():
            sheet = workbook.create_sheet(name)
            for row in sheet_rows:
                sheet.append(row)
        workbook.save(buffer)
        workbook.close()
    else:
        workbook = xlwt.Workbook()
        for name, sheet_rows in sheets.items():
            sheet = workbook.add_sheet(name)
            for row_number, row in enumerate(sheet_rows):
                for column, value in enumerate(row):
                    if value is not None:
                        sheet.write(row_number, column, value)
        workbook.save(buffer)
    return SimpleUploadedFile(f"tags.{extension}", buffer.getvalue())


class TagFileParsingTests(SimpleTestCase):
    def test_xls_and_xlsx_support_repeated_keys_blank_rows_and_numeric_values(self):
        rows = [
            ["标签", "标签值"],
            [" 产品名称 ", " 产品A "],
            [None, None],
            ["产品名称", "产品B"],
            ["版本号", 0],
            ["版本号", 1],
        ]
        for extension in ("xls", "xlsx"):
            with self.subTest(extension=extension):
                self.assertEqual(
                    parse_tag_file(tag_file(rows, extension)),
                    [
                        {"key": "产品名称", "value": "产品A"},
                        {"key": "产品名称", "value": "产品B"},
                        {"key": "版本号", "value": "0"},
                        {"key": "版本号", "value": "1"},
                    ],
                )

    def test_supplied_template_is_importable(self):
        tags = parse_tag_file(SimpleUploadedFile("标签模板.xlsx", TAG_TEMPLATE_PATH.read_bytes()))
        self.assertEqual(len(tags), 8)
        self.assertEqual(tags[0], {"key": "产品名称", "value": "产品A"})
        self.assertEqual(tags[-1], {"key": "版本类型", "value": "专业版"})

    def test_all_nonempty_sheets_are_validated(self):
        for extension in ("xls", "xlsx"):
            with self.subTest(extension=extension):
                tags = parse_tag_file(
                    tag_file(
                        [["标签", "标签值"], ["产品", "A"]],
                        extension,
                        {"第二页": [["标签", "标签值"], ["产品", "B"]], "空白页": []},
                    )
                )
                self.assertEqual(tags, [{"key": "产品", "value": "A"}, {"key": "产品", "value": "B"}])

    def test_missing_cells_and_length_limits_report_sheet_and_row(self):
        invalid_rows = ([None, "值"], ["标签", None], [" ", "值"], ["a" * 65, "值"], ["标签", "a" * 129])
        for extension in ("xls", "xlsx"):
            for row in invalid_rows:
                with self.subTest(extension=extension, row=row):
                    with self.assertRaises(AppApiException) as raised:
                        parse_tag_file(tag_file([["标签", "标签值"], row], extension))
                    self.assertIn("Sheet1", str(raised.exception.message))
                    self.assertIn("第 2 行", str(raised.exception.message))

    def test_wrong_headers_and_extra_columns_are_rejected(self):
        invalid_rows = (
            [["标签值", "标签"], ["A", "B"]],
            [["key", "value"], ["A", "B"]],
            [["标签", "标签值", "其他"], ["A", "B", "C"]],
            [["标签", "标签值"], ["A", "B", "C"]],
            [[None, None], ["A", "B"]],
        )
        for extension in ("xls", "xlsx"):
            for rows in invalid_rows:
                with self.subTest(extension=extension, rows=rows):
                    with self.assertRaises(AppApiException):
                        parse_tag_file(tag_file(rows, extension))

    def test_empty_and_header_only_workbooks_are_rejected(self):
        for extension in ("xls", "xlsx"):
            for rows in ([], [["标签", "标签值"]], [["标签", "标签值"], [None, None]]):
                with self.subTest(extension=extension, rows=rows):
                    with self.assertRaises(AppApiException) as raised:
                        parse_tag_file(tag_file(rows, extension))
                    self.assertIn("没有可导入", str(raised.exception.message))

    def test_unsupported_and_corrupt_files_are_rejected(self):
        for name in ("tags.csv", "tags.xlsx", "tags.xls"):
            with self.subTest(name=name):
                with self.assertRaises(AppApiException):
                    parse_tag_file(SimpleUploadedFile(name, b"invalid workbook"))

    def test_uppercase_extension_is_supported(self):
        file = tag_file([["标签", "标签值"], ["产品", "A"]])
        file.name = "tags.XLSX"
        self.assertEqual(parse_tag_file(file), [{"key": "产品", "value": "A"}])


class TagImportSerializerTests(SimpleTestCase):
    def serializer(self, workspace_id="default"):
        return TagImportSerializers(data={"workspace_id": workspace_id, "knowledge_id": KNOWLEDGE_ID})

    def import_tags(self, serializer, instance):
        # File validation and insertion run normally; unit tests do not open a database transaction.
        return TagImportSerializers.import_tags.__wrapped__(serializer, instance)

    def test_existing_and_file_duplicate_pairs_are_skipped(self):
        file = tag_file([["标签", "标签值"], ["产品", "A"], ["产品", "B"], ["产品", "B"]])
        with (
            patch("knowledge.serializers.tag_import.QuerySet") as scope_query,
            patch("knowledge.serializers.tag.QuerySet") as tag_query,
            patch("knowledge.serializers.tag.Tag.objects.bulk_create") as bulk_create,
        ):
            scope_query.return_value.filter.return_value.exists.return_value = True
            tag_query.return_value.filter.return_value.exists.return_value = True
            tag_query.return_value.filter.return_value.values_list.return_value = [("产品", "A")]
            self.assertIsNone(self.import_tags(self.serializer(), {"file": file}))
            tags = bulk_create.call_args.args[0]
            self.assertEqual([(tag.key, tag.value) for tag in tags], [("产品", "B")])
            self.assertEqual(str(tags[0].knowledge_id), KNOWLEDGE_ID)

    def test_reimport_with_only_existing_tags_does_not_insert(self):
        with (
            patch("knowledge.serializers.tag_import.QuerySet") as scope_query,
            patch("knowledge.serializers.tag.QuerySet") as tag_query,
            patch("knowledge.serializers.tag.Tag.objects.bulk_create") as bulk_create,
        ):
            scope_query.return_value.filter.return_value.exists.return_value = True
            tag_query.return_value.filter.return_value.exists.return_value = True
            tag_query.return_value.filter.return_value.values_list.return_value = [("产品", "A")]
            self.import_tags(self.serializer(), {"file": tag_file([["标签", "标签值"], ["产品", "A"]])})
            bulk_create.assert_not_called()

    def test_invalid_later_row_or_sheet_prevents_all_inserts(self):
        for rows, extra_sheets in (
            ([["标签", "标签值"], ["产品", "A"], ["产品", None]], None),
            ([["标签", "标签值"], ["产品", "A"]], {"坏表": [["错误表头"]]}),
        ):
            with self.subTest(rows=rows, extra_sheets=extra_sheets):
                with (
                    patch("knowledge.serializers.tag_import.QuerySet") as scope_query,
                    patch("knowledge.serializers.tag_import.TagSerializers.Create") as create,
                ):
                    scope_query.return_value.filter.return_value.exists.return_value = True
                    with self.assertRaises(AppApiException):
                        self.import_tags(self.serializer(), {"file": tag_file(rows, extra_sheets=extra_sheets)})
                    create.assert_not_called()

    def test_cross_workspace_and_missing_knowledge_are_rejected_before_parsing(self):
        for workspace_id in ("other", "None"):
            for operation in ("import", "export"):
                with self.subTest(workspace_id=workspace_id, operation=operation):
                    with (
                        patch("knowledge.serializers.tag_import.QuerySet") as scope_query,
                        patch("knowledge.serializers.tag_import.parse_tag_file") as parse,
                    ):
                        scope_query.return_value.filter.return_value.exists.return_value = False
                        with self.assertRaises(AppApiException):
                            if operation == "import":
                                self.import_tags(self.serializer(workspace_id), {})
                            else:
                                self.serializer(workspace_id).export_template()
                        scope_query.return_value.filter.assert_called_once_with(
                            id=UUID(KNOWLEDGE_ID), workspace_id=workspace_id
                        )
                        parse.assert_not_called()

    def test_missing_or_empty_upload_is_rejected(self):
        for data in ({}, {"file": SimpleUploadedFile("tags.xlsx", b"")}):
            with self.subTest(data=data):
                with patch("knowledge.serializers.tag_import.QuerySet") as scope_query:
                    scope_query.return_value.filter.return_value.exists.return_value = True
                    with self.assertRaises(ValidationError):
                        self.import_tags(self.serializer(), data)

    def test_download_returns_the_original_template(self):
        with patch("knowledge.serializers.tag_import.QuerySet") as scope_query:
            scope_query.return_value.filter.return_value.exists.return_value = True
            response = self.serializer().export_template()
        self.assertEqual(response.content, TAG_TEMPLATE_PATH.read_bytes())
        self.assertEqual(response["Content-Type"], TAG_TEMPLATE_CONTENT_TYPE)
        self.assertIn(quote("标签模板.xlsx"), response["Content-Disposition"])


class TagImportEndpointTests(SimpleTestCase):
    prefix = f"/workspace/default/knowledge/{KNOWLEDGE_ID}/tags"

    def test_routes_and_request_schema(self):
        self.assertIs(
            resolve(f"{self.prefix}/import", urlconf="knowledge.urls").func.view_class, KnowledgeTagImportView
        )
        self.assertIs(
            resolve(f"{self.prefix}/template/export", urlconf="knowledge.urls").func.view_class,
            KnowledgeTagTemplateView,
        )
        request_schema = TagImportAPI.get_request()["multipart/form-data"]
        self.assertEqual(request_schema["required"], ["file"])
        self.assertEqual(request_schema["properties"]["file"]["format"], "binary")
        self.assertEqual(
            [parameter.name for parameter in TagImportAPI.get_parameters()], ["workspace_id", "knowledge_id"]
        )
        self.assertIn((200, TAG_TEMPLATE_CONTENT_TYPE), TagTemplateExportAPI.get_response())
        self.assertEqual(KnowledgeTagImportView.authentication_classes, [TokenAuth])
        self.assertEqual(KnowledgeTagTemplateView.authentication_classes, [TokenAuth])

    def test_generated_schema_supports_binary_upload_and_template_download(self):
        route = "workspace/<str:workspace_id>/knowledge/<str:knowledge_id>/tags"
        schema = SchemaGenerator(
            patterns=[
                path(f"{route}/import", KnowledgeTagImportView.as_view()),
                path(f"{route}/template/export", KnowledgeTagTemplateView.as_view()),
            ]
        ).get_schema(public=True)
        prefix = "/workspace/{workspace_id}/knowledge/{knowledge_id}/tags"
        upload_schema = schema["paths"][f"{prefix}/import"]["post"]["requestBody"]["content"]["multipart/form-data"][
            "schema"
        ]
        self.assertEqual(upload_schema["properties"]["file"]["format"], "binary")
        response = schema["paths"][f"{prefix}/template/export"]["get"]["responses"]["200"]
        self.assertEqual(response["content"][TAG_TEMPLATE_CONTENT_TYPE]["schema"]["format"], "binary")

    def test_multipart_import_passes_the_file_and_scope_to_serializer(self):
        file = tag_file([["标签", "标签值"], ["产品", "A"]])
        request = APIRequestFactory().post(f"{self.prefix}/import", {"file": file}, format="multipart")
        force_authenticate(request, user=SimpleNamespace(id=USER_ID, type="TEST"))
        with (
            patch("common.auth.authentication._build") as permissions,
            patch("common.log.log.Log"),
            patch("knowledge.views.tag_import.get_knowledge_operation_object", return_value={}),
            patch("knowledge.views.tag_import.TagImportSerializers") as serializer,
        ):
            permissions.return_value.hasPermission.return_value = True
            serializer.return_value.import_tags.return_value = None
            response = KnowledgeTagImportView.as_view()(request, workspace_id="default", knowledge_id=KNOWLEDGE_ID)
            payload = json.loads(response.content)
            self.assertEqual(payload["code"], 200)
            self.assertIsNone(payload["data"])
            serializer.assert_called_once_with(data={"workspace_id": "default", "knowledge_id": KNOWLEDGE_ID})
            uploaded = serializer.return_value.import_tags.call_args.args[0]["file"]
            self.assertEqual(uploaded.name, "tags.xlsx")
            self.assertEqual(parse_tag_file(uploaded), [{"key": "产品", "value": "A"}])
            scope = {"workspace_id": "default", "knowledge_id": KNOWLEDGE_ID}
            self.assertEqual(
                permissions.call_args.args[0][0](request, scope),
                PermissionConstants.KNOWLEDGE_TAG_CREATE.get_workspace_knowledge_permission()(request, scope),
            )

    def test_permissions_prevent_import_and_template_download(self):
        for view, method in ((KnowledgeTagImportView, "post"), (KnowledgeTagTemplateView, "get")):
            with self.subTest(view=view.__name__):
                with (
                    patch("common.auth.authentication._build") as permissions,
                    patch("knowledge.views.tag_import.TagImportSerializers") as serializer,
                ):
                    permissions.return_value.hasPermission.return_value = False
                    with self.assertRaises(AppUnauthorizedFailed):
                        getattr(view(), method)(SimpleNamespace(), workspace_id="default", knowledge_id=KNOWLEDGE_ID)
                    serializer.assert_not_called()
